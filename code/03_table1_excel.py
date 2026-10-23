"""
ECON 20 Project - Table 1 formatted like the course example (panels A/B, SD in brackets).
Reads final_data.dta (made by 01_build_data.do). If that file does not exist yet,
it rebuilds the same sample from raw/ with the same rules as 01_build_data.do.
Output: output/Table1_means_formatted.xlsx
"""
import os, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

os.chdir(os.path.dirname(os.path.abspath(__file__)) + "/..")

def rebuild():
    m = pd.read_stata("raw/compustat_annual_main.dta"); s = pd.read_stata("raw/compustat_annual_csho.dta")
    m = m[(m.indfmt == 'INDL') & (m['loc'] == 'USA') & (m.state != '')]
    m = m[~m.duplicated(['gvkey', 'fyear'], keep=False)]
    s = s[s.indfmt == 'INDL']; s = s[~s.duplicated(['gvkey', 'fyear'], keep=False)][['gvkey', 'fyear', 'csho', 'apdedate', 'gsector']]
    d = m.merge(s, on=['gvkey', 'fyear'], how='left')
    d = d[d.exchg.isin([11, 12, 14]) & ~d.sic.str.startswith('67')]
    fyeyr = np.where(d.fyr >= 6, d.fyear, d.fyear + 1)
    alt = pd.to_datetime(dict(year=fyeyr, month=d.fyr.fillna(12).astype(int), day=1), errors='coerce') + pd.offsets.MonthEnd(0)
    d['fye'] = d.apdedate.fillna(alt)
    d['post'] = (d.fye >= '2018-10-01').astype(int); d['ca'] = (d.state == 'CA').astype(int)
    d['tx'] = (d.state == 'TX').astype(int); d['wa'] = (d.state == 'WA').astype(int)
    d['mktcap'] = d.prcc_f * d.csho; d['lnmktcap'] = np.log(d.mktcap.where(d.mktcap > 0))
    A = d['at'].where(d['at'] > 0)
    d['lnassets'] = np.log(A); d['roa'] = d.ni / A; d['leverage'] = d['lt'] / A
    d['multinational'] = ((d.pifo != 0) & d.pifo.notna()).astype(int)
    d['age'] = (d.fye.dt.year - d.ipodate.dt.year).where(lambda x: x >= 0)
    d['nyse'] = (d.exchg == 11).astype(int); d['tech'] = (d.gsector == '45').astype(int)
    for v in ['roa', 'leverage']:
        lo, hi = d[v].quantile([.01, .99]); d[v] = d[v].clip(lo, hi)
    b = pd.read_stata("raw/boardex_orgsummary.dta"); b = b[b.hocountryname == 'United States']
    b = b.drop_duplicates(['companyid', 'annualreportdate'])[['companyid', 'annualreportdate', 'genderratio', 'numberdirectors']]
    b['yr'] = b.annualreportdate.dt.year
    b = b.sort_values(['companyid', 'yr', 'annualreportdate']).drop_duplicates(['companyid', 'yr'], keep='last')
    l = pd.read_stata("raw/link_boardex_crsp_compustat.dta")
    l = l[l.preferred == 1].sort_values(['companyid', 'score']).drop_duplicates('companyid')[['companyid', 'GVKEY']].rename(columns={'GVKEY': 'gvkey'})
    b = b.merge(l, on='companyid').sort_values(['gvkey', 'yr', 'annualreportdate']).drop_duplicates(['gvkey', 'yr'], keep='last')
    b['women'] = (b.numberdirectors * (1 - b.genderratio)).round(); b['zerowomen'] = (b.women == 0).astype(int)
    b = b.rename(columns={'numberdirectors': 'boardsize'})
    d['yr'] = d.fye.dt.year
    d = d.merge(b[['gvkey', 'yr', 'boardsize', 'women', 'zerowomen']], on=['gvkey', 'yr'], how='left')
    d = d[(d.fye >= '2010-01-01') & (d.fye <= '2025-12-31')]
    d['mainwindow'] = (d.fye <= '2019-12-31').astype(int)
    return d[d.wa == 0]

if os.path.exists("final_data.dta"):
    d = pd.read_stata("final_data.dta"); source = "final_data.dta (Stata)"
else:
    d = rebuild(); source = "rebuilt from raw/ in Python (run 01_build_data.do to replace)"
d = d[(d.mainwindow == 1) & (d.wa == 0)]

rows = [  # (variable, label line 1, label line 2, number format for mean)
    ('lnmktcap', 'Ln(Market Cap)', None, '0.000'),
    ('mktcap', 'Market Cap', '($ millions)', '#,##0'),
    ('lnassets', 'Ln(Total Assets)', None, '0.000'),
    ('roa', 'Net Income /', 'Total Assets', '0.000'),
    ('leverage', 'Total Liabilities /', 'Total Assets', '0.000'),
    ('multinational', 'Multinational', '(foreign income)', '0.000'),
    ('age', 'Years Since IPO', None, '0.00'),
    ('nyse', 'Listed on NYSE', None, '0.000'),
    ('tech', 'Tech Sector', '(GICS 45)', '0.000'),
    ('boardsize', 'Board Size', None, '0.00'),
    ('women', 'Women on Board', None, '0.000'),
    ('zerowomen', 'No Women', 'on Board', '0.000'),
]
cols = [('California', lambda x: x.ca == 1), ('All Other States', lambda x: x.ca == 0), ('Texas', lambda x: x.tx == 1)]

thin = Side(style='thin'); F = 'Calibri'
wb = Workbook(); ws = wb.active; ws.title = 'Table 1'
ws.column_dimensions['A'].width = 24
for c in 'BCD': ws.column_dimensions[c].width = 16
def put(r, c, v, **k):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=F, size=12, bold=k.get('bold', False), italic=k.get('italic', False), underline=k.get('underline'))
    cell.alignment = Alignment(horizontal=k.get('h', 'center'), vertical='center', wrap_text=k.get('wrap', False))
    if 'fmt' in k: cell.number_format = k['fmt']
    return cell

put(1, 1, 'Table 1. Means by Headquarters State, Before and After California SB 826', bold=True)
ws.merge_cells('A1:D1')
put(3, 2, 'Treatment', underline='single')
put(3, 3, 'Comparison', underline='single'); ws.merge_cells('C3:D3')
for j, (name, _) in enumerate(cols): put(4, 2 + j, name)
for c in range(1, 5): ws.cell(row=4, column=c).border = Border(bottom=thin)

r = 6
panels = [('A. Before the Law (fiscal years ending Jan 2010 - Sep 2018)', 0),
          ('B. After the Law (fiscal years ending Oct 2018 - Dec 2019)', 1)]
for title, p in panels:
    put(r, 1, title, italic=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); r += 1
    sub = d[d.post == p]
    for var, l1, l2, fmt in rows:
        put(r, 1, l1, h='left'); put(r + 1, 1, ('  ' + l2) if l2 else None, h='left')
        for j, (_, f) in enumerate(cols):
            x = sub[f(sub)][var].dropna()
            put(r, 2 + j, float(x.mean()), fmt=fmt)
            put(r + 1, 2 + j, float(x.std()), fmt='"["' + fmt + '"]"')
        r += 2
    put(r, 1, 'Firm-Years', h='left')
    put(r + 1, 1, 'Firms', h='left')
    for j, (_, f) in enumerate(cols):
        x = sub[f(sub)]
        put(r, 2 + j, int(len(x)), fmt='#,##0'); put(r + 1, 2 + j, int(x.gvkey.nunique()), fmt='#,##0')
    r += 3
for c in range(1, 5): ws.cell(row=r - 1, column=c).border = Border(bottom=thin)

note = ("Data Source: Compustat North America Fundamentals Annual and BoardEx North America, "
        "accessed through WRDS. Sample: US-headquartered firms listed on NYSE, AMEX or Nasdaq, "
        "excluding funds and shell companies (SIC 67xx); fiscal years ending Jan 2010 - Dec 2019. "
        "Washington firms are excluded (Washington passed its own board-diversity law in 2020); "
        "'All Other States' includes Texas. Headquarters state is the current state reported by Compustat. "
        "Board variables are available for firm-years matched to BoardEx (about 80%). "
        "Net income/assets and liabilities/assets are winsorized at the 1st and 99th percentiles. "
        "Unweighted means; standard deviations in brackets.")
put(r, 1, note, h='left', wrap=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws.row_dimensions[r].height = 120
ws.cell(row=r, column=1).font = Font(name=F, size=11)
ws.cell(row=r, column=1).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
ws.cell(row=r + 2, column=1, value='Numbers computed from: ' + source).font = Font(name=F, size=9, italic=True, color='808080')
ws.sheet_view.showGridLines = False
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
wb.save('output/Table1_means_formatted.xlsx')
print('saved; source =', source, '; rows used =', len(d))
